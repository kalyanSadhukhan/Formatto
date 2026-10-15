import hashlib
import logging
import os
import re
import xml.etree.ElementTree as ET
from typing import Any, Dict, List, Optional, Set, Tuple

import docx
from docx import Document
from docx.oxml.ns import nsmap

logger = logging.getLogger(__name__)

# Ensure namespaces
if "v" not in nsmap:
    nsmap["v"] = "urn:schemas-microsoft-com:vml"
if "mc" not in nsmap:
    nsmap["mc"] = "http://schemas.openxmlformats.org/markup-compatibility/2006"
if "m" not in nsmap:
    nsmap["m"] = "http://schemas.openxmlformats.org/officeDocument/2006/math"
if "wp" not in nsmap:
    nsmap["wp"] = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_V = "urn:schemas-microsoft-com:vml"
NS_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"

NS = {
    "w": NS_W,
    "r": NS_R,
    "mc": NS_MC,
    "v": NS_V,
    "m": NS_M,
    "wp": NS_WP,
}

# Conversion factors:
# 1 pt = 20 twips (dxa)
# 1 half-point = 0.5 pt (used by w:sz)
TWIPS_TO_PT = 1.0 / 20.0
HALF_PT_TO_PT = 0.5


def _xp(elem, query: str) -> list:
    """Safe xpath evaluation handling both BaseOxmlElement and raw lxml elements."""
    if elem is None:
        return []
    try:
        return elem.xpath(query, namespaces=NS)
    except TypeError:
        return elem.xpath(query)


class TemplateExtractor:
    """
    Deterministic DOCX Template Extractor (Specification V2, Phase 3A & 3B).

    Extracts document defaults, paragraph and character styles (including inheritance chains,
    raw vs inherited properties, and property lineages), outline levels, font attributes,
    paragraph spacing/indentation, numbering definitions, table formatting, page setup/sections
    (dimensions, margins, orientation), headers, footers, page-number fields, unsupported
    features, and direct formatting overrides without using an LLM.

    Produces structured `TemplateProfile` and a normalized list of `TemplateConstraint` objects
    with explicit provenance tracking (EXTRACTED, CONFIGURED, INFERRED, UNSUPPORTED) and
    strict enforceability distinction (inferred properties are never marked enforceable).
    """

    def __init__(self, file_path: str, template_id: Optional[str] = None):
        self.file_path = file_path
        if not os.path.exists(file_path):
            raise ValueError(f"Template file not found: {file_path}")

        try:
            with open(file_path, "rb") as f:
                self.file_bytes = f.read()
            self.template_hash = template_id or hashlib.sha256(self.file_bytes).hexdigest()[:8]
        except Exception as e:
            raise ValueError(f"Failed to read file bytes: {e}")

        try:
            self.doc = Document(file_path)
        except Exception as e:
            raise ValueError(f"Failed to open DOCX template: {e}")

        self.doc_defaults: Dict[str, Any] = {}
        self.styles_map: Dict[str, Any] = {}
        self.unsupported_features: List[Dict[str, Any]] = []

    def _twips_to_pt(self, val: Any) -> Optional[float]:
        if val is None:
            return None
        try:
            return round(float(val) * TWIPS_TO_PT, 2)
        except (ValueError, TypeError):
            return None

    def _half_pt_to_pt(self, val: Any) -> Optional[float]:
        if val is None:
            return None
        try:
            return round(float(val) * HALF_PT_TO_PT, 2)
        except (ValueError, TypeError):
            return None

    def _extract_doc_defaults(self) -> Dict[str, Any]:
        """Extract w:docDefaults from styles.xml."""
        defaults: Dict[str, Any] = {
            "font_name": "Calibri",  # Word global fallback
            "font_size_pt": 11.0,
            "line_spacing": 1.15,
            "space_after_pt": 8.0,
            "provenance": "CONFIGURED",
        }

        doc_defs_nodes = _xp(self.doc.styles.element, "./w:docDefaults")
        if not doc_defs_nodes:
            return defaults

        doc_defs = doc_defs_nodes[0]
        # Font defaults from rPrDefault
        r_fonts = _xp(doc_defs, "./w:rPrDefault/w:rPr/w:rFonts/@w:ascii")
        if r_fonts:
            defaults["font_name"] = r_fonts[0]
            defaults["provenance"] = "EXTRACTED"
        else:
            # Check for theme font
            theme_fonts = _xp(doc_defs, "./w:rPrDefault/w:rPr/w:rFonts/@w:asciiTheme")
            if theme_fonts:
                defaults["font_theme"] = theme_fonts[0]

        sz_nodes = _xp(doc_defs, "./w:rPrDefault/w:rPr/w:sz/@w:val")
        if sz_nodes:
            defaults["font_size_pt"] = self._half_pt_to_pt(sz_nodes[0])
            defaults["provenance"] = "EXTRACTED"

        # Paragraph defaults from pPrDefault
        spacing_nodes = _xp(doc_defs, "./w:pPrDefault/w:pPr/w:spacing")
        if spacing_nodes:
            sp = spacing_nodes[0]
            after = sp.attrib.get(f"{{{NS_W}}}after")
            if after:
                defaults["space_after_pt"] = self._twips_to_pt(after)
                defaults["provenance"] = "EXTRACTED"
            line = sp.attrib.get(f"{{{NS_W}}}line")
            line_rule = sp.attrib.get(f"{{{NS_W}}}lineRule", "auto")
            if line and line_rule == "auto":
                try:
                    defaults["line_spacing"] = round(float(line) / 240.0, 2)
                    defaults["provenance"] = "EXTRACTED"
                except ValueError:
                    pass

        return defaults

    def _parse_style_properties(self, style_elem) -> Tuple[Dict[str, Any], List[str]]:
        """
        Parse raw properties explicitly declared by a <w:style> element.
        Also returns list of undetermined property keys.
        """
        props: Dict[str, Any] = {}
        undetermined: List[str] = []

        # 1. Font properties (rPr)
        r_fonts = _xp(style_elem, "./w:rPr/w:rFonts/@w:ascii")
        if r_fonts:
            props["font_name"] = r_fonts[0]
        else:
            ascii_theme = _xp(style_elem, "./w:rPr/w:rFonts/@w:asciiTheme")
            if ascii_theme:
                undetermined.append(f"font_name_theme_dependent:{ascii_theme[0]}")

        sz = _xp(style_elem, "./w:rPr/w:sz/@w:val")
        if sz:
            props["font_size_pt"] = self._half_pt_to_pt(sz[0])

        b_nodes = _xp(style_elem, "./w:rPr/w:b")
        if b_nodes:
            val = b_nodes[0].attrib.get(f"{{{NS_W}}}val")
            props["bold"] = False if val in ("0", "false") else True

        i_nodes = _xp(style_elem, "./w:rPr/w:i")
        if i_nodes:
            val = i_nodes[0].attrib.get(f"{{{NS_W}}}val")
            props["italic"] = False if val in ("0", "false") else True

        u_nodes = _xp(style_elem, "./w:rPr/w:u/@w:val")
        if u_nodes:
            props["underline"] = u_nodes[0] != "none"

        color_nodes = _xp(style_elem, "./w:rPr/w:color/@w:val")
        if color_nodes:
            props["color"] = color_nodes[0]

        # 2. Paragraph properties (pPr)
        jc = _xp(style_elem, "./w:pPr/w:jc/@w:val")
        if jc:
            props["alignment"] = jc[0].lower()

        outline_lvl = _xp(style_elem, "./w:pPr/w:outlineLvl/@w:val")
        if outline_lvl:
            try:
                props["outline_level"] = int(outline_lvl[0]) + 1
            except ValueError:
                pass

        spacing = _xp(style_elem, "./w:pPr/w:spacing")
        if spacing:
            sp = spacing[0]
            before = sp.attrib.get(f"{{{NS_W}}}before")
            if before:
                props["space_before_pt"] = self._twips_to_pt(before)
            after = sp.attrib.get(f"{{{NS_W}}}after")
            if after:
                props["space_after_pt"] = self._twips_to_pt(after)
            line = sp.attrib.get(f"{{{NS_W}}}line")
            line_rule = sp.attrib.get(f"{{{NS_W}}}lineRule", "auto")
            props["line_spacing_rule"] = line_rule
            if line:
                if line_rule == "auto":
                    try:
                        props["line_spacing"] = round(float(line) / 240.0, 2)
                    except ValueError:
                        pass
                else:
                    props["line_spacing"] = self._twips_to_pt(line)

        ind = _xp(style_elem, "./w:pPr/w:ind")
        if ind:
            in_elem = ind[0]
            left = in_elem.attrib.get(f"{{{NS_W}}}left")
            if left:
                props["left_indent_pt"] = self._twips_to_pt(left)
            right = in_elem.attrib.get(f"{{{NS_W}}}right")
            if right:
                props["right_indent_pt"] = self._twips_to_pt(right)
            first_line = in_elem.attrib.get(f"{{{NS_W}}}firstLine")
            if first_line:
                props["first_line_indent_pt"] = self._twips_to_pt(first_line)
            hanging = in_elem.attrib.get(f"{{{NS_W}}}hanging")
            if hanging:
                props["hanging_indent_pt"] = self._twips_to_pt(hanging)

        # Check for unsupported framePr / drop cap
        if _xp(style_elem, "./w:pPr/w:framePr"):
            undetermined.append("frame_drop_cap")

        return props, undetermined

    def _extract_styles(self) -> Dict[str, Dict[str, Any]]:
        """
        Extract all styles, resolve inheritance (basedOn) chains, and distinguish:
        - raw_declared_properties (direct properties declared by the style)
        - inherited_properties (properties inherited from ancestor styles)
        - property_lineage (lineage path showing where each property was defined/overridden)
        - resolved_effective_properties (effective merged values)
        - undetermined_properties (theme fonts, drop caps, etc.)
        """
        raw_styles: Dict[str, Dict[str, Any]] = {}

        for s_elem in _xp(self.doc.styles.element, "./w:style"):
            sid = s_elem.attrib.get(f"{{{NS_W}}}styleId")
            if not sid:
                continue

            stype = s_elem.attrib.get(f"{{{NS_W}}}type", "paragraph")
            name_nodes = _xp(s_elem, "./w:name/@w:val")
            name = name_nodes[0] if name_nodes else sid

            based_on_nodes = _xp(s_elem, "./w:basedOn/@w:val")
            based_on = based_on_nodes[0] if based_on_nodes else None

            next_style_nodes = _xp(s_elem, "./w:next/@w:val")
            next_style = next_style_nodes[0] if next_style_nodes else None

            is_custom = s_elem.attrib.get(f"{{{NS_W}}}customStyle") == "1"

            direct_props, undetermined = self._parse_style_properties(s_elem)

            # Outline level heuristic by name if not in pPr
            if "outline_level" not in direct_props:
                name_lower = name.lower()
                if name_lower.startswith("heading"):
                    try:
                        lvl_num = int(name_lower.replace("heading", "").strip())
                        direct_props["outline_level"] = lvl_num
                    except ValueError:
                        pass
                elif name_lower == "title":
                    direct_props["outline_level"] = 1
                elif name_lower == "subtitle":
                    direct_props["outline_level"] = 2

            raw_styles[sid] = {
                "style_id": sid,
                "style_name": name,
                "style_type": stype,
                "based_on": based_on,
                "next_style": next_style,
                "is_custom": is_custom,
                "direct_properties": direct_props,
                "raw_declared_properties": direct_props.copy(),
                "undetermined_properties": undetermined,
            }

        # Resolve inheritance for all styles
        resolved_styles: Dict[str, Dict[str, Any]] = {}
        for sid, sdata in raw_styles.items():
            chain: List[str] = [sid]
            curr = sdata.get("based_on")
            visited: Set[str] = {sid}

            while curr and curr in raw_styles and curr not in visited:
                chain.append(curr)
                visited.add(curr)
                curr = raw_styles[curr].get("based_on")

            # Resolve property lineage and effective values
            # Chain is ordered: [leaf, parent, grandparent, ...]
            # Ancestors in reverse order: [root_ancestor, ..., parent, leaf]
            resolved_props: Dict[str, Any] = {}
            inherited_props: Dict[str, Any] = {}
            property_lineage: Dict[str, List[str]] = {}

            # Seed paragraph styles with document defaults if applicable
            if sdata["style_type"] == "paragraph":
                for def_k in ["font_name", "font_size_pt", "space_after_pt", "line_spacing"]:
                    if def_k in self.doc_defaults:
                        resolved_props[def_k] = self.doc_defaults[def_k]
                        property_lineage[def_k] = ["docDefaults"]

            for ancestor_id in reversed(chain):
                ancestor_direct = raw_styles[ancestor_id]["direct_properties"]
                for k, v in ancestor_direct.items():
                    resolved_props[k] = v
                    if k in property_lineage:
                        property_lineage[k].append(ancestor_id)
                    else:
                        property_lineage[k] = [ancestor_id]

            # Distinguish inherited properties: properties in resolved_props not directly declared in leaf
            leaf_direct = sdata["direct_properties"]
            for k, v in resolved_props.items():
                if k not in leaf_direct:
                    inherited_props[k] = v

            resolved_styles[sid] = {
                **sdata,
                "raw_declared_properties": leaf_direct,
                "inherited_properties": inherited_props,
                "property_lineage": property_lineage,
                "resolved_effective_properties": resolved_props,
                "undetermined_properties": sdata["undetermined_properties"],
                # Backward compatibility
                "direct_properties": leaf_direct,
                "resolved_properties": resolved_props,
                "property_provenance": {k: lineage[-1] for k, lineage in property_lineage.items()},
                "inheritance_chain": chain,
            }

        return resolved_styles

    def _extract_numbering(self) -> Dict[str, Any]:
        """Extract numbering schemes from numbering.xml."""
        numbering_schemes: Dict[str, Any] = {}
        try:
            num_part = self.doc.part.numbering_part
            if not num_part:
                return numbering_schemes

            root = ET.fromstring(num_part._element.xml)

            abstract_map: Dict[str, Any] = {}
            for abstract in root.findall(f".//{{{NS_W}}}abstractNum"):
                aid = abstract.attrib.get(f"{{{NS_W}}}abstractNumId")
                if not aid:
                    continue
                levels: Dict[int, Any] = {}
                for lvl in abstract.findall(f"./{{{NS_W}}}lvl"):
                    ilvl = int(lvl.attrib.get(f"{{{NS_W}}}ilvl", "0"))
                    num_fmt = lvl.find(f"./{{{NS_W}}}numFmt")
                    fmt_val = num_fmt.attrib.get(f"{{{NS_W}}}val", "decimal") if num_fmt is not None else "decimal"
                    lvl_text = lvl.find(f"./{{{NS_W}}}lvlText")
                    txt_val = lvl_text.attrib.get(f"{{{NS_W}}}val", "") if lvl_text is not None else ""
                    start_node = lvl.find(f"./{{{NS_W}}}start")
                    start_val = int(start_node.attrib.get(f"{{{NS_W}}}val", "1")) if start_node is not None else 1

                    ind_node = lvl.find(f"./{{{NS_W}}}pPr/{{{NS_W}}}ind")
                    left_ind = self._twips_to_pt(ind_node.attrib.get(f"{{{NS_W}}}left")) if ind_node is not None else None
                    hanging_ind = self._twips_to_pt(ind_node.attrib.get(f"{{{NS_W}}}hanging")) if ind_node is not None else None

                    levels[ilvl] = {
                        "level": ilvl,
                        "num_fmt": fmt_val,
                        "lvl_text": txt_val,
                        "start": start_val,
                        "left_indent_pt": left_ind,
                        "hanging_indent_pt": hanging_ind,
                    }
                abstract_map[aid] = levels

            for num in root.findall(f".//{{{NS_W}}}num"):
                nid = num.attrib.get(f"{{{NS_W}}}numId")
                abstract_ref = num.find(f"./{{{NS_W}}}abstractNumId")
                aid = abstract_ref.attrib.get(f"{{{NS_W}}}val") if abstract_ref is not None else None
                if nid and aid and aid in abstract_map:
                    numbering_schemes[nid] = {
                        "num_id": nid,
                        "abstract_num_id": aid,
                        "levels": abstract_map[aid],
                    }

        except Exception as e:
            logger.debug(f"Note: Could not parse template numbering: {e}")

        return numbering_schemes

    def _classify_page_size(self, width_pt: float, height_pt: float) -> str:
        """Classify standard page sizes (A4, Letter, Legal, or Custom)."""
        w, h = min(width_pt, height_pt), max(width_pt, height_pt)
        if abs(w - 595.3) < 4.0 and abs(h - 841.9) < 4.0:
            return "A4"
        elif abs(w - 612.0) < 4.0 and abs(h - 792.0) < 4.0:
            return "Letter"
        elif abs(w - 612.0) < 4.0 and abs(h - 1008.0) < 4.0:
            return "Legal"
        return "Custom"

    def _extract_sections(self) -> List[Dict[str, Any]]:
        """
        Extract section layout properties, margins, headers, footers, and page numbers.
        
        Reads OpenXML elements directly from sectPr and relationship parts to avoid
        mutational side effects of python-docx (which fabricates synthetic header/footer
        references if paragraphs on an unlinked header are accessed).
        """
        sections_data: List[Dict[str, Any]] = []

        for s_idx, section in enumerate(self.doc.sections):
            sect_elm = section._sectPr

            # 1. Page dimensions and orientation directly from <w:pgSz>
            pg_sz_nodes = _xp(sect_elm, "./w:pgSz")
            if pg_sz_nodes:
                w_attr = pg_sz_nodes[0].attrib.get(f"{{{NS_W}}}w")
                h_attr = pg_sz_nodes[0].attrib.get(f"{{{NS_W}}}h")
                orient_attr = pg_sz_nodes[0].attrib.get(f"{{{NS_W}}}orient", "portrait")
                w_pt = self._twips_to_pt(w_attr) if w_attr else round(section.page_width.pt, 2)
                h_pt = self._twips_to_pt(h_attr) if h_attr else round(section.page_height.pt, 2)
                orient = orient_attr.lower()
            else:
                w_pt = round(section.page_width.pt, 2)
                h_pt = round(section.page_height.pt, 2)
                orient = "portrait" if section.orientation == 0 else "landscape"

            size_name = self._classify_page_size(w_pt, h_pt)

            type_nodes = _xp(sect_elm, "./w:type/@w:val")
            break_type = type_nodes[0] if type_nodes else "nextPage"

            # 2. Margins directly from <w:pgMar>
            pg_mar_nodes = _xp(sect_elm, "./w:pgMar")
            if pg_mar_nodes:
                top_attr = pg_mar_nodes[0].attrib.get(f"{{{NS_W}}}top")
                bot_attr = pg_mar_nodes[0].attrib.get(f"{{{NS_W}}}bottom")
                left_attr = pg_mar_nodes[0].attrib.get(f"{{{NS_W}}}left")
                right_attr = pg_mar_nodes[0].attrib.get(f"{{{NS_W}}}right")
                hdr_attr = pg_mar_nodes[0].attrib.get(f"{{{NS_W}}}header")
                ftr_attr = pg_mar_nodes[0].attrib.get(f"{{{NS_W}}}footer")
                margins = {
                    "top": self._twips_to_pt(top_attr) if top_attr else round(section.top_margin.pt, 2),
                    "bottom": self._twips_to_pt(bot_attr) if bot_attr else round(section.bottom_margin.pt, 2),
                    "left": self._twips_to_pt(left_attr) if left_attr else round(section.left_margin.pt, 2),
                    "right": self._twips_to_pt(right_attr) if right_attr else round(section.right_margin.pt, 2),
                    "header": self._twips_to_pt(hdr_attr) if hdr_attr else 36.0,
                    "footer": self._twips_to_pt(ftr_attr) if ftr_attr else 36.0,
                }
            else:
                margins = {
                    "top": round(section.top_margin.pt, 2),
                    "bottom": round(section.bottom_margin.pt, 2),
                    "left": round(section.left_margin.pt, 2),
                    "right": round(section.right_margin.pt, 2),
                    "header": round(section.header_distance.pt, 2) if section.header_distance else 36.0,
                    "footer": round(section.footer_distance.pt, 2) if section.footer_distance else 36.0,
                }

            # 3. Columns
            col_nodes = _xp(sect_elm, "./w:cols/@w:num")
            columns_count = int(col_nodes[0]) if col_nodes else 1

            # 4. Headers & Footers: read relationship IDs BEFORE touching python-docx objects
            hdr_refs = _xp(sect_elm, "./w:headerReference")
            ftr_refs = _xp(sect_elm, "./w:footerReference")

            hdr_rids = [r.attrib.get(f"{{{NS_R}}}id") for r in hdr_refs if r.attrib.get(f"{{{NS_R}}}id")]
            ftr_rids = [r.attrib.get(f"{{{NS_R}}}id") for r in ftr_refs if r.attrib.get(f"{{{NS_R}}}id")]

            # Section linking logic:
            # Section 0 is never linked to previous.
            # Section k > 0 is linked if no explicit reference element is present in sectPr.
            if s_idx == 0:
                is_header_linked = False
                is_footer_linked = False
                has_hdr = bool(hdr_rids)
                has_ftr = bool(ftr_rids)
            else:
                is_header_linked = (len(hdr_rids) == 0)
                is_footer_linked = (len(ftr_rids) == 0)
                has_hdr = bool(hdr_rids) or (is_header_linked and sections_data[s_idx - 1]["has_header"])
                has_ftr = bool(ftr_rids) or (is_footer_linked and sections_data[s_idx - 1]["has_footer"])

            # Read text and check page numbers directly from related parts
            page_field_xpath = (
                './/w:fldSimple[contains(@w:instr, "PAGE")] '
                '| .//w:instrText[contains(text(), "PAGE")]'
            )

            hdr_text = None
            has_page_num_hdr = False
            for rid in hdr_rids:
                if rid in self.doc.part.related_parts:
                    part_elem = self.doc.part.related_parts[rid].element
                    t_nodes = _xp(part_elem, ".//w:t/text()")
                    if t_nodes:
                        hdr_text = " ".join(t.strip() for t in t_nodes if t.strip())
                    if _xp(part_elem, page_field_xpath):
                        has_page_num_hdr = True

            ftr_text = None
            has_page_num_ftr = False
            for rid in ftr_rids:
                if rid in self.doc.part.related_parts:
                    part_elem = self.doc.part.related_parts[rid].element
                    t_nodes = _xp(part_elem, ".//w:t/text()")
                    if t_nodes:
                        ftr_text = " ".join(t.strip() for t in t_nodes if t.strip())
                    if _xp(part_elem, page_field_xpath):
                        has_page_num_ftr = True

            # If linked to previous section, inherit text and page number detection
            if is_header_linked and s_idx > 0:
                hdr_text = sections_data[s_idx - 1]["header_text"]
                has_page_num_hdr = sections_data[s_idx - 1].get("_has_page_num_hdr", False)
            if is_footer_linked and s_idx > 0:
                ftr_text = sections_data[s_idx - 1]["footer_text"]
                has_page_num_ftr = sections_data[s_idx - 1].get("_has_page_num_ftr", False)

            has_page_number = has_page_num_hdr or has_page_num_ftr

            sections_data.append({
                "section_index": s_idx,
                "break_type": break_type,
                "page_width_pt": w_pt,
                "page_height_pt": h_pt,
                "orientation": orient,
                "page_size_name": size_name,
                "margins_pt": margins,
                "columns_count": columns_count,
                "has_header": has_hdr,
                "has_footer": has_ftr,
                "is_header_linked": is_header_linked,
                "is_footer_linked": is_footer_linked,
                "has_page_number": has_page_number,
                "header_text": hdr_text,
                "footer_text": ftr_text,
                "header_rids": hdr_rids,
                "footer_rids": ftr_rids,
                "_has_page_num_hdr": has_page_num_hdr,
                "_has_page_num_ftr": has_page_num_ftr,
            })

        return sections_data

    def _extract_table_styles(self) -> Dict[str, Any]:
        """Extract table styles and border/cell margin definitions."""
        table_styles: Dict[str, Any] = {}

        for s_elem in _xp(self.doc.styles.element, "./w:style[@w:type='table']"):
            sid = s_elem.attrib.get(f"{{{NS_W}}}styleId")
            name_nodes = _xp(s_elem, "./w:name/@w:val")
            name = name_nodes[0] if name_nodes else sid

            borders = _xp(s_elem, "./w:tblPr/w:tblBorders")
            has_borders = False
            border_details = {}
            if borders:
                has_borders = True
                for b_side in ["top", "bottom", "left", "right", "insideH", "insideV"]:
                    side_node = borders[0].find(f"./{{{NS_W}}}{b_side}")
                    if side_node is not None:
                        val = side_node.attrib.get(f"{{{NS_W}}}val", "none")
                        sz = side_node.attrib.get(f"{{{NS_W}}}sz")
                        color = side_node.attrib.get(f"{{{NS_W}}}color")
                        border_details[b_side] = {
                            "val": val,
                            "sz": float(sz) / 8.0 if sz else None,
                            "color": color,
                        }

            cell_mar = _xp(s_elem, "./w:tblPr/w:tblCellMar")
            cell_margins_pt = {}
            if cell_mar:
                for side in ["top", "bottom", "left", "right"]:
                    m_node = cell_mar[0].find(f"./{{{NS_W}}}{side}")
                    if m_node is not None:
                        w_val = m_node.attrib.get(f"{{{NS_W}}}w")
                        cell_margins_pt[side] = self._twips_to_pt(w_val)

            table_styles[sid] = {
                "style_id": sid,
                "style_name": name,
                "has_borders": has_borders,
                "border_details": border_details,
                "cell_margins_pt": cell_margins_pt,
            }

        return table_styles

    def _detect_direct_formatting_overrides(self) -> List[Dict[str, Any]]:
        """Inspect sample paragraphs in the template body to detect direct formatting overrides."""
        overrides: List[Dict[str, Any]] = []
        for p_idx, p in enumerate(self.doc.paragraphs):
            style_id = p.style.style_id if p.style else "Normal"
            p_elem = p._p

            pPr = p_elem.pPr
            p_overrides: Dict[str, Any] = {}
            if pPr is not None:
                jc = _xp(pPr, "./w:jc/@w:val")
                if jc:
                    p_overrides["alignment"] = jc[0].lower()
                sp = _xp(pPr, "./w:spacing")
                if sp:
                    after = sp[0].attrib.get(f"{{{NS_W}}}after")
                    if after:
                        p_overrides["space_after_pt"] = self._twips_to_pt(after)
                ind = _xp(pPr, "./w:ind")
                if ind:
                    left = ind[0].attrib.get(f"{{{NS_W}}}left")
                    if left:
                        p_overrides["left_indent_pt"] = self._twips_to_pt(left)

            r_overrides: Dict[str, Any] = {}
            for r in p.runs:
                rPr = r._r.rPr
                if rPr is not None:
                    b = _xp(rPr, "./w:b")
                    if b:
                        r_overrides["bold"] = True
                    i = _xp(rPr, "./w:i")
                    if i:
                        r_overrides["italic"] = True
                    sz = _xp(rPr, "./w:sz/@w:val")
                    if sz:
                        r_overrides["font_size_pt"] = self._half_pt_to_pt(sz[0])
                    color = _xp(rPr, "./w:color/@w:val")
                    if color:
                        r_overrides["color"] = color[0]

            if p_overrides or r_overrides:
                overrides.append({
                    "paragraph_index": p_idx,
                    "style_id": style_id,
                    "paragraph_overrides": p_overrides,
                    "run_overrides": r_overrides,
                    "text_snippet": p.text[:50].strip() if p.text else None,
                })
        return overrides

    def _detect_unsupported_features(self) -> List[Dict[str, Any]]:
        """Detect formatting features present in template that cannot be deterministically evaluated."""
        unsupported: List[Dict[str, Any]] = []

        # 1. Floating drawing objects (wp:anchor)
        anchors = _xp(self.doc.element.body, ".//w:drawing//wp:anchor")
        if anchors:
            unsupported.append({
                "feature": "floating_anchored_drawing",
                "count": len(anchors),
                "description": "Floating drawings with absolute page coordinates are unsupported in flow layout.",
            })

        # 2. Legacy VML objects
        vml = _xp(self.doc.element.body, ".//v:shape | .//v:group")
        if vml:
            unsupported.append({
                "feature": "legacy_vml_shape",
                "count": len(vml),
                "description": "Legacy VML vector shapes are deprecated in ISO/IEC 29500 and unsupported.",
            })

        # 3. Word MathML
        math = _xp(self.doc.element.body, ".//m:oMath | .//m:oMathPara")
        if math:
            unsupported.append({
                "feature": "word_math_equations",
                "count": len(math),
                "description": "Office MathML equations require specialized typography engine.",
            })

        # 4. Embedded OLE objects
        ole = _xp(self.doc.element.body, ".//w:object")
        if ole:
            unsupported.append({
                "feature": "embedded_ole_object",
                "count": len(ole),
                "description": "Embedded OLE objects/ActiveX controls are unsupported.",
            })

        # 5. Theme-dependent fonts in styles without theme resolution
        theme_fonts = _xp(self.doc.styles.element, ".//w:rFonts[@w:asciiTheme and not(@w:ascii)]")
        if theme_fonts:
            unsupported.append({
                "feature": "theme_dependent_fonts",
                "count": len(theme_fonts),
                "description": "Fonts referencing theme slots (e.g. minorHAnsi) without static font name.",
            })

        return unsupported

    def _construct_constraints(
        self,
        sections: List[Dict[str, Any]],
        styles: Dict[str, Dict[str, Any]],
        defaults: Dict[str, Any],
        unsupported: List[Dict[str, Any]],
    ) -> List[Dict[str, Any]]:
        """
        Construct formal TemplateConstraint instances adhering strictly to docs/03_DATA_CONTRACTS.md.
        
        Provenance Contract:
        - EXTRACTED: Explicitly declared in XML of this section/style. (enforceable = True, confidence = 1.0)
        - INFERRED: Derived via inheritance (basedOn) or section linking. (enforceable = False, confidence = 0.85)
          * Inferred properties are never marked enforceable.
        - CONFIGURED: Document-level or platform fallback values. (enforceable = False, confidence = 0.5)
        - UNSUPPORTED: Feature present in document but cannot be evaluated. (enforceable = False, confidence = 0.0)

        Constraints are sorted deterministically by (category, target, property).
        """
        constraints: List[Dict[str, Any]] = []

        # 1. Document Defaults (target: "document.defaults")
        if defaults.get("font_name"):
            is_extracted = defaults.get("provenance") == "EXTRACTED"
            constraints.append({
                "id": "tc_doc_default_font_name",
                "category": "typography",
                "target": "document.defaults",
                "property": "font_name",
                "expected_value": defaults["font_name"],
                "provenance": defaults.get("provenance", "CONFIGURED"),
                "confidence": 1.0 if is_extracted else 0.5,
                "enforceable": is_extracted,
                "evidence": {"default_type": "w:rPrDefault", "provenance": defaults.get("provenance")},
            })

        if defaults.get("font_size_pt"):
            is_extracted = defaults.get("provenance") == "EXTRACTED"
            constraints.append({
                "id": "tc_doc_default_font_size_pt",
                "category": "typography",
                "target": "document.defaults",
                "property": "font_size_pt",
                "expected_value": defaults["font_size_pt"],
                "provenance": defaults.get("provenance", "CONFIGURED"),
                "confidence": 1.0 if is_extracted else 0.5,
                "enforceable": is_extracted,
                "evidence": {"default_type": "w:rPrDefault", "provenance": defaults.get("provenance")},
            })

        if defaults.get("space_after_pt") is not None:
            is_extracted = defaults.get("provenance") == "EXTRACTED"
            constraints.append({
                "id": "tc_doc_default_space_after_pt",
                "category": "paragraph",
                "target": "document.defaults",
                "property": "space_after_pt",
                "expected_value": defaults["space_after_pt"],
                "provenance": defaults.get("provenance", "CONFIGURED"),
                "confidence": 1.0 if is_extracted else 0.5,
                "enforceable": is_extracted,
                "evidence": {"default_type": "w:pPrDefault", "provenance": defaults.get("provenance")},
            })

        if defaults.get("line_spacing") is not None:
            is_extracted = defaults.get("provenance") == "EXTRACTED"
            constraints.append({
                "id": "tc_doc_default_line_spacing",
                "category": "paragraph",
                "target": "document.defaults",
                "property": "line_spacing",
                "expected_value": defaults["line_spacing"],
                "provenance": defaults.get("provenance", "CONFIGURED"),
                "confidence": 1.0 if is_extracted else 0.5,
                "enforceable": is_extracted,
                "evidence": {"default_type": "w:pPrDefault", "provenance": defaults.get("provenance")},
            })

        # 2. Section & Page constraints (category: "page" and "header_footer")
        for s in sections:
            s_idx = s["section_index"]
            target = f"section.{s_idx}"

            # Page size & orientation
            constraints.append({
                "id": f"tc_sec_{s_idx}_page_size",
                "category": "page",
                "target": target,
                "property": "page_size_name",
                "expected_value": s["page_size_name"],
                "provenance": "EXTRACTED",
                "confidence": 1.0,
                "enforceable": True,
                "evidence": {"width_pt": s["page_width_pt"], "height_pt": s["page_height_pt"]},
            })
            constraints.append({
                "id": f"tc_sec_{s_idx}_orientation",
                "category": "page",
                "target": target,
                "property": "orientation",
                "expected_value": s["orientation"],
                "provenance": "EXTRACTED",
                "confidence": 1.0,
                "enforceable": True,
                "evidence": {"w:orient": s["orientation"]},
            })

            # Margins
            for side in ["top", "bottom", "left", "right"]:
                val = s["margins_pt"].get(side)
                if val is not None:
                    constraints.append({
                        "id": f"tc_sec_{s_idx}_margin_{side}",
                        "category": "page",
                        "target": target,
                        "property": f"margin_{side}_pt",
                        "expected_value": val,
                        "provenance": "EXTRACTED",
                        "confidence": 1.0,
                        "enforceable": True,
                        "evidence": {"side": side, "pt": val},
                    })

            # Header / Footer & Page Numbering
            is_linked = s.get("is_footer_linked", False)
            constraints.append({
                "id": f"tc_sec_{s_idx}_has_page_number",
                "category": "header_footer",
                "target": target,
                "property": "has_page_number",
                "expected_value": s["has_page_number"],
                "provenance": "INFERRED" if is_linked else "EXTRACTED",
                "confidence": 0.85 if is_linked else 1.0,
                "enforceable": False if is_linked else True,
                "evidence": {
                    "field": "PAGE",
                    "present": s["has_page_number"],
                    "is_footer_linked": is_linked,
                    "explicit_ftr_rids": s.get("footer_rids", []),
                },
            })

        # 3. Style constraints for key roles (category: "typography", "paragraph", "structure")
        key_style_roles = {
            "Normal": "style.Normal",
            "BodyText": "style.BodyText",
            "Heading1": "style.Heading1",
            "Heading2": "style.Heading2",
            "Heading3": "style.Heading3",
            "Title": "style.Title",
            "Subtitle": "style.Subtitle",
            "Caption": "style.Caption",
            "ListParagraph": "style.ListParagraph",
        }

        for sid, target in key_style_roles.items():
            sdata = styles.get(sid)
            if not sdata:
                # Search by style_name
                sdata = next(
                    (s for s in styles.values() if s["style_name"].replace(" ", "").lower() == sid.lower()),
                    None
                )
            if not sdata:
                continue

            r_props = sdata.get("resolved_effective_properties", sdata.get("resolved_properties", {}))
            d_props = sdata.get("raw_declared_properties", sdata.get("direct_properties", {}))
            lineage_map = sdata.get("property_lineage", {})

            # Typography: font_name
            font_name = r_props.get("font_name")
            if font_name and ("font_name" in d_props or "font_name" in lineage_map):
                is_direct = "font_name" in d_props
                constraints.append({
                    "id": f"tc_{sid}_font_name",
                    "category": "typography",
                    "target": target,
                    "property": "font_name",
                    "expected_value": font_name,
                    "provenance": "EXTRACTED" if is_direct else "INFERRED",
                    "confidence": 1.0 if is_direct else 0.85,
                    "enforceable": True if is_direct else False,
                    "evidence": {
                        "direct": is_direct,
                        "lineage": lineage_map.get("font_name", []),
                    },
                })

            # Typography: font_size_pt
            font_size = r_props.get("font_size_pt")
            if font_size and ("font_size_pt" in d_props or "font_size_pt" in lineage_map):
                is_direct = "font_size_pt" in d_props
                constraints.append({
                    "id": f"tc_{sid}_font_size_pt",
                    "category": "typography",
                    "target": target,
                    "property": "font_size_pt",
                    "expected_value": font_size,
                    "provenance": "EXTRACTED" if is_direct else "INFERRED",
                    "confidence": 1.0 if is_direct else 0.85,
                    "enforceable": True if is_direct else False,
                    "evidence": {
                        "direct": is_direct,
                        "lineage": lineage_map.get("font_size_pt", []),
                    },
                })

            # Typography: bold
            bold = r_props.get("bold")
            if bold is not None and ("bold" in d_props or "bold" in lineage_map):
                is_direct = "bold" in d_props
                constraints.append({
                    "id": f"tc_{sid}_bold",
                    "category": "typography",
                    "target": target,
                    "property": "bold",
                    "expected_value": bold,
                    "provenance": "EXTRACTED" if is_direct else "INFERRED",
                    "confidence": 1.0 if is_direct else 0.85,
                    "enforceable": True if is_direct else False,
                    "evidence": {
                        "direct": is_direct,
                        "lineage": lineage_map.get("bold", []),
                    },
                })

            # Paragraph: alignment
            alignment = r_props.get("alignment")
            if alignment and ("alignment" in d_props or "alignment" in lineage_map):
                is_direct = "alignment" in d_props
                constraints.append({
                    "id": f"tc_{sid}_alignment",
                    "category": "paragraph",
                    "target": target,
                    "property": "alignment",
                    "expected_value": alignment,
                    "provenance": "EXTRACTED" if is_direct else "INFERRED",
                    "confidence": 1.0 if is_direct else 0.85,
                    "enforceable": True if is_direct else False,
                    "evidence": {
                        "direct": is_direct,
                        "lineage": lineage_map.get("alignment", []),
                    },
                })

            # Paragraph: line_spacing
            line_sp = r_props.get("line_spacing")
            if line_sp is not None and ("line_spacing" in d_props or "line_spacing" in lineage_map):
                is_direct = "line_spacing" in d_props
                constraints.append({
                    "id": f"tc_{sid}_line_spacing",
                    "category": "paragraph",
                    "target": target,
                    "property": "line_spacing",
                    "expected_value": line_sp,
                    "provenance": "EXTRACTED" if is_direct else "INFERRED",
                    "confidence": 1.0 if is_direct else 0.85,
                    "enforceable": True if is_direct else False,
                    "evidence": {
                        "direct": is_direct,
                        "lineage": lineage_map.get("line_spacing", []),
                    },
                })

            # Paragraph: space_after_pt
            space_after = r_props.get("space_after_pt")
            if space_after is not None and ("space_after_pt" in d_props or "space_after_pt" in lineage_map):
                is_direct = "space_after_pt" in d_props
                constraints.append({
                    "id": f"tc_{sid}_space_after_pt",
                    "category": "paragraph",
                    "target": target,
                    "property": "space_after_pt",
                    "expected_value": space_after,
                    "provenance": "EXTRACTED" if is_direct else "INFERRED",
                    "confidence": 1.0 if is_direct else 0.85,
                    "enforceable": True if is_direct else False,
                    "evidence": {
                        "direct": is_direct,
                        "lineage": lineage_map.get("space_after_pt", []),
                    },
                })

            # Structure: outline_level
            outline = r_props.get("outline_level")
            if outline:
                is_direct = "outline_level" in d_props
                constraints.append({
                    "id": f"tc_{sid}_outline_level",
                    "category": "structure",
                    "target": target,
                    "property": "outline_level",
                    "expected_value": outline,
                    "provenance": "EXTRACTED" if is_direct else "INFERRED",
                    "confidence": 1.0 if is_direct else 0.85,
                    "enforceable": True if is_direct else False,
                    "evidence": {
                        "direct": is_direct,
                        "outline_level": outline,
                        "lineage": lineage_map.get("outline_level", []),
                    },
                })

        # 4. Unsupported Features
        for idx, u in enumerate(unsupported):
            constraints.append({
                "id": f"tc_unsupported_{idx:02d}_{u['feature']}",
                "category": "object",
                "target": f"unsupported.{u['feature']}",
                "property": "rendered_fidelity",
                "expected_value": "UNSUPPORTED",
                "provenance": "UNSUPPORTED",
                "confidence": 0.0,
                "enforceable": False,
                "evidence": u,
            })

        # Deterministic sort by category, target, property
        constraints.sort(key=lambda c: (c["category"], c["target"], c["property"]))
        return constraints

    def _extract_template_tables(self) -> Tuple[List[Dict[str, Any]], bool]:
        """
        Extract pre-existing tables from the template document to identify
        structural/form templates that must have their structure preserved.
        """
        table_structs: List[Dict[str, Any]] = []
        is_structured = False

        form_keywords = [
            "project", "batch", "id", "date", "signature", "item", "sl. no.",
            "examiner", "guide", "student", "attendance", "present", "absent",
            "checklist", "review", "consent", "declaration"
        ]

        for idx, tbl in enumerate(self.doc.tables):
            rows_cnt = len(tbl.rows)
            cols_cnt = len(tbl.columns)
            headers: List[str] = []
            if tbl.rows:
                headers = [c.text.strip().replace("\n", " ") for c in tbl.rows[0].cells]

            # Check for merged cells (multiple cells pointing to same _tc element)
            has_merges = False
            tc_seen = set()
            for r in tbl.rows:
                for c in r.cells:
                    if c._tc in tc_seen:
                        has_merges = True
                        break
                    tc_seen.add(c._tc)
                if has_merges:
                    break

            # Check if this table has form characteristics
            is_form = False
            all_text_lower = " ".join(headers).lower()
            if any(k in all_text_lower for k in form_keywords):
                is_form = True

            # Collect signature labels
            sig_labels: List[str] = []
            for h in headers:
                if "signature" in h.lower() or "sign." in h.lower():
                    sig_labels.append(h)

            table_structs.append({
                "table_index": idx,
                "rows_count": rows_cnt,
                "columns_count": cols_cnt,
                "header_texts": headers,
                "has_merged_cells": has_merges,
                "is_form_table": is_form,
                "signature_labels": sig_labels,
            })

        # A template is structured if it has form characteristics (e.g. blank entry rows for manual logging,
        # form field labels like Project/Batch ID, Attendance, Examiner signatures, Consent) and at least 2 tables.
        form_markers = [
            "project/batch id", "attendance", "guide signature", "examiner 1 signature",
            "task performed / work description", "project submission checklist",
            "student consent form", "capstone project-i diary", "daily attendance record"
        ]
        doc_text = " ".join(p.text.lower() for p in self.doc.paragraphs[:50])
        for t in table_structs:
            doc_text += " " + " ".join(h.lower() for h in t["header_texts"])

        keyword_matches = sum(1 for k in form_markers if k in doc_text)
        blank_rows = 0
        for tbl in self.doc.tables:
            for r in tbl.rows[1:]:
                if not "".join(c.text.strip() for c in r.cells):
                    blank_rows += 1

        if (keyword_matches >= 2 or blank_rows >= 5) and len(self.doc.tables) >= 2:
            is_structured = True

        return table_structs, is_structured


    def extract_template(self) -> Dict[str, Any]:
        """
        Execute deterministic template extraction and return a comprehensive profile dict.
        """
        self.doc_defaults = self._extract_doc_defaults()
        styles = self._extract_styles()
        sections = self._extract_sections()
        numbering = self._extract_numbering()
        tables = self._extract_table_styles()
        unsupported = self._detect_unsupported_features()
        overrides = self._detect_direct_formatting_overrides()
        template_tables, is_structured = self._extract_template_tables()

        constraints_dicts = self._construct_constraints(
            sections=sections,
            styles=styles,
            defaults=self.doc_defaults,
            unsupported=unsupported,
        )

        return {
            "template_id": self.template_hash,
            "file_path": self.file_path,
            "document_defaults": self.doc_defaults,
            "sections": sections,
            "styles": styles,
            "numbering_schemes": numbering,
            "table_styles": tables,
            "template_tables": template_tables,
            "is_structured_template": is_structured,
            "constraints": constraints_dicts,
            "unsupported_features": unsupported,
            "direct_formatting_overrides": overrides,
        }

