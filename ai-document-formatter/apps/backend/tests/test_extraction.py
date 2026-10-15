import pytest
import os
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches
from app.extraction.docx_extractor import DocxExtractor

@pytest.fixture
def temp_docx(tmp_path):
    doc_path = tmp_path / "test_doc.docx"
    doc = Document()
    
    # 1. Simple Paragraph
    doc.add_paragraph("This is a simple paragraph.")
    
    # 2. Headings with multiple levels
    doc.add_heading("Heading 1", level=1)
    doc.add_heading("Heading 2", level=2)
    
    # 3. Mixed paragraph/heading
    doc.add_paragraph("Paragraph after heading.")
    
    # 4. Bullet list
    doc.add_paragraph("Bullet item 1", style='List Bullet')
    
    # 5. Numbered list
    doc.add_paragraph("Number item 1", style='List Number')
    
    # 6. Table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    table.cell(1, 0).text = "A2"
    table.cell(1, 1).text = "B2"
    
    # 7. Images
    # We will simulate an image by adding an empty picture (if we have one) 
    # or we just rely on a paragraph with w:drawing manually injected, 
    # but for simplicity we'll add a picture using python-docx if possible.
    # Since we don't have an image file, let's create a blank image file first.
    img_path = tmp_path / "test.png"
    with open(img_path, "wb") as f:
        # A tiny valid PNG
        f.write(b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82')
    doc.add_picture(str(img_path), width=Inches(1.0))
    
    # 8. Captions
    # Create caption style if not exists
    styles = doc.styles
    if 'Caption' not in styles:
        styles.add_style('Caption', WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("Figure 1: Test image", style='Caption')
    
    # 9. Mixed ordering
    doc.add_paragraph("Trailing paragraph.")
    
    # 13. Very large text element
    large_text = "A" * 10000
    doc.add_paragraph(large_text)
    
    # 14. Unicode text
    doc.add_paragraph("Unicode test: 你好 🌍")
    
    doc.save(doc_path)
    return str(doc_path)

def test_extract_deterministic_order(temp_docx):
    extractor = DocxExtractor(temp_docx)
    elements = extractor.extract()
    
    # Check ordering and types
    assert len(elements) >= 12
    
    assert elements[0].type == "paragraph"
    assert elements[0].text == "This is a simple paragraph."
    
    assert elements[1].type == "heading"
    assert elements[1].level == 1
    assert elements[1].text == "Heading 1"
    
    assert elements[2].type == "heading"
    assert elements[2].level == 2
    assert elements[2].text == "Heading 2"
    
    assert elements[3].type == "paragraph"
    assert elements[3].text == "Paragraph after heading."
    
    assert elements[4].type == "bullet_list"
    assert elements[4].text == "Bullet item 1"
    
    assert elements[5].type == "numbered_list"
    assert elements[5].text == "Number item 1"
    
    assert elements[6].type == "table"
    assert elements[6].metadata["row_count"] == 2
    assert elements[6].metadata["column_count"] == 2
    assert elements[6].metadata["cell_text"][0][0] == "A1"
    
    assert elements[7].type == "figure"
    assert elements[7].metadata["image_count"] >= 1
    
    assert elements[8].type == "caption"
    assert elements[8].text == "Figure 1: Test image"
    
    assert elements[9].type == "paragraph"
    assert elements[9].text == "Trailing paragraph."

    # Very large text
    assert elements[10].type == "paragraph"
    assert len(elements[10].text) == 10000
    
    # Unicode text
    assert elements[11].type == "paragraph"
    assert "你好 🌍" in elements[11].text
    
    # Check stable IDs and source order
    for i, el in enumerate(elements):
        assert el.id.startswith("el_")
        assert el.source_order == i + 1

def test_extract_empty_doc(tmp_path):
    doc_path = tmp_path / "empty.docx"
    doc = Document()
    doc.save(doc_path)
    extractor = DocxExtractor(str(doc_path))
    elements = extractor.extract()
    assert len(elements) == 0

def test_extract_invalid_doc():
    with pytest.raises(ValueError):
        DocxExtractor("non_existent_file.docx")
