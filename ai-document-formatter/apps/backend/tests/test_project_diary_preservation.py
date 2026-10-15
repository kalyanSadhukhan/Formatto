import os
import tempfile
import pytest
import docx

from app.pipeline.orchestrator import DocumentFormattingPipeline
from app.domain.schemas.data_contracts import JobState

FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "fixtures", "project_diary")
TEMPLATE_PATH = os.path.join(FIXTURES_DIR, "annexure_iii_template.docx")
SOURCE_PATH = os.path.join(FIXTURES_DIR, "team_093_raw_source.docx")

FLOW_FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "fixtures", "flow_report")
FLOW_SOURCE_PATH = os.path.join(FLOW_FIXTURES_DIR, "sample.docx")
if not os.path.exists(FLOW_SOURCE_PATH):
    FLOW_SOURCE_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(__file__)))), "sample.docx")

FLOW_TEMPLATE_PATH = os.path.join(FLOW_FIXTURES_DIR, "reference_template.docx")
if not os.path.exists(FLOW_TEMPLATE_PATH):
    FLOW_TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(__file__)))), "reference_template.docx")


@pytest.fixture(scope="module")
def pipeline_execution():
    """Execute end-to-end pipeline on Project Diary fixtures once for all module tests."""
    out_dir = tempfile.mkdtemp(prefix="test_diary_exec_")
    pipeline = DocumentFormattingPipeline(output_dir=out_dir)
    res = pipeline.run(
        source_path=SOURCE_PATH,
        template_path=TEMPLATE_PATH,
        job_id="test_diary_regress",
    )
    assert res.status == JobState.COMPLETED
    assert res.output_docx_path is not None
    assert os.path.exists(res.output_docx_path)

    out_doc = docx.Document(res.output_docx_path)
    tpl_doc = docx.Document(TEMPLATE_PATH)
    return {
        "result": res,
        "out_doc": out_doc,
        "tpl_doc": tpl_doc,
    }


def test_section_boundaries_preservation(pipeline_execution):
    """Test 8: Preservation of the required section boundaries."""
    out_doc = pipeline_execution["out_doc"]
    tpl_doc = pipeline_execution["tpl_doc"]

    assert len(out_doc.sections) == 9
    assert len(out_doc.sections) == len(tpl_doc.sections)

    # Verify start types
    assert out_doc.sections[0].start_type == tpl_doc.sections[0].start_type
    for i in range(1, 9):
        assert out_doc.sections[i].start_type == tpl_doc.sections[i].start_type


def test_cover_page_metadata_table_preservation_and_population(pipeline_execution):
    """Test 1: Preservation of the cover-page metadata table and its fields."""
    out_doc = pipeline_execution["out_doc"]
    assert len(out_doc.tables) >= 1

    t0 = out_doc.tables[0]
    assert len(t0.rows) == 10
    assert len(t0.columns) == 4

    # Verify field keys and populated values
    assert "Project/Batch ID" in t0.rows[0].cells[0].text
    assert t0.rows[0].cells[1].text == "ASAC-CP-093"

    assert "Project Guide Name" in t0.rows[1].cells[0].text
    assert t0.rows[1].cells[1].text == "Dr. Nishant Tripathi"

    assert "Title of the Project" in t0.rows[2].cells[0].text
    assert "Formatto" in t0.rows[2].cells[1].text

    assert "Domain" in t0.rows[3].cells[0].text
    assert "Document Intelligence" in t0.rows[3].cells[1].text

    assert "Team Leader Name" in t0.rows[4].cells[0].text
    assert t0.rows[4].cells[1].text == "Kalyan Sadhukhan"


def test_student_details_table_preservation_and_population(pipeline_execution):
    """Test 2: Preservation and population of the student details sub-rows in Table 0."""
    out_doc = pipeline_execution["out_doc"]
    t0 = out_doc.tables[0]

    expected_students = [
        ("Anirudhh Bhaskar N J", "2023BIFT07AED004", "B.Tech IT (CC)"),
        ("Keerthana K N", "2023BIFT07AED057", "B.Tech IT (DA)"),
        ("Kalyan Sadhukhan", "2023BCSE07AED217", "B.Tech CSE (AIML)"),
        ("Pichikala Kali Bhargav Sai", "2023BCSE07AED278", "B.Tech CSE (AIML)"),
    ]

    for idx, (name, reg, branch) in enumerate(expected_students):
        row = t0.rows[6 + idx]
        assert row.cells[1].text == name
        assert row.cells[2].text == reg
        assert row.cells[3].text == branch


def test_attendance_table_preservation(pipeline_execution):
    """Test 3: Preservation of all required attendance tables."""
    out_doc = pipeline_execution["out_doc"]
    assert len(out_doc.tables) >= 2

    t1 = out_doc.tables[1]
    assert len(t1.rows) == 26
    assert len(t1.columns) == 5

    headers = [c.text.strip() for c in t1.rows[0].cells]
    assert "Date" in headers[0]
    assert "Day" in headers[1]
    assert "Task Performed / Work Description" in headers[2]
    assert "Guide Signature" in headers[4]


def test_blank_attendance_rows_preservation(pipeline_execution):
    """Test 4: Preservation of blank attendance rows (not replaced by paragraphs)."""
    out_doc = pipeline_execution["out_doc"]
    t1 = out_doc.tables[1]

    # Verify that all 25 rows under the header are blank for physical logging
    blank_rows = 0
    for r in t1.rows[1:]:
        content = "".join(c.text.strip() for c in r.cells)
        if not content:
            blank_rows += 1

    assert blank_rows == 25


def test_weekly_meeting_form_preservation_and_population(pipeline_execution):
    """Test 5: Preservation and population of the weekly meeting form."""
    out_doc = pipeline_execution["out_doc"]
    full_text = "\n".join(p.text for p in out_doc.paragraphs)

    assert "Week No: 01" in full_text
    assert "Anirudhh Bhaskar N J" in full_text
    assert "2023BIFT07AED004" in full_text
    assert "☐ Present" in full_text
    assert "OF 211" in full_text


def test_review_evaluation_forms_preservation(pipeline_execution):
    """Test 6: Preservation of review evaluation forms and checklists."""
    out_doc = pipeline_execution["out_doc"]

    # Table 2: Review 1 evaluation signatures
    t2 = out_doc.tables[2]
    assert len(t2.rows) == 3
    assert len(t2.columns) == 3
    assert "Project Guide Signature" in t2.rows[0].cells[0].text
    assert "Examiner 1 Signature" in t2.rows[0].cells[1].text
    assert "Examiner 2 Signature" in t2.rows[0].cells[2].text

    # Table 3: Review 2 evaluation signatures
    t3 = out_doc.tables[3]
    assert len(t3.rows) == 3
    assert len(t3.columns) == 3

    # Table 4: Project Submission Checklist
    t4 = out_doc.tables[4]
    assert len(t4.rows) == 10
    assert len(t4.columns) == 4
    assert "Item" in t4.rows[0].cells[0].text
    assert "Project Report Submitted" in t4.rows[1].cells[0].text
    assert "☐" in t4.rows[1].cells[1].text


def test_student_consent_form_preservation(pipeline_execution):
    """Test 7: Preservation and population of the consent form."""
    out_doc = pipeline_execution["out_doc"]
    full_text = "\n".join(p.text for p in out_doc.paragraphs)

    assert "Academic Year: 2026-2027" in full_text
    assert "Semester: VII" in full_text
    assert "Degree Program: B. Tech" in full_text
    assert "Associate Professor" in full_text
    assert "Declaration by the Students" in full_text

    # Table 5: Student Details in Consent Form
    t5 = out_doc.tables[5]
    assert len(t5.rows) == 5
    assert len(t5.columns) == 3
    assert t5.rows[1].cells[1].text == "Anirudhh Bhaskar N J"
    assert t5.rows[1].cells[2].text == "2023BIFT07AED004"

    # Table 6: Student Signatures in Consent Form
    t6 = out_doc.tables[6]
    assert len(t6.rows) == 5
    assert len(t6.columns) == 3
    assert t6.rows[1].cells[1].text == "Anirudhh Bhaskar N J"
    assert t6.rows[1].cells[2].text == ""  # Must remain blank for manual physical signing


def test_headers_and_footers_preservation(pipeline_execution):
    """Test 9: Preservation of headers, footers, and institutional layouts."""
    out_doc = pipeline_execution["out_doc"]
    tpl_doc = pipeline_execution["tpl_doc"]

    for i in range(len(tpl_doc.sections)):
        tpl_hdr = " ".join(p.text.strip() for p in tpl_doc.sections[i].header.paragraphs if p.text.strip())
        out_hdr = " ".join(p.text.strip() for p in out_doc.sections[i].header.paragraphs if p.text.strip())
        assert out_hdr == tpl_hdr


def test_correct_population_of_supported_fields(pipeline_execution):
    """Test 10: Correct population of supported fields."""
    res = pipeline_execution["result"]
    assert res.manifest is not None
    mapped_targets = {m.ir_target_id for m in res.manifest.mappings if m.status == "mapped"}

    assert "project_batch_id" in mapped_targets
    assert "project_guide_name" in mapped_targets
    assert "project_title" in mapped_targets
    assert "domain" in mapped_targets
    assert "team_leader_name" in mapped_targets
    assert "student_1" in mapped_targets
    assert "student_2" in mapped_targets
    assert "student_3" in mapped_targets
    assert "student_4" in mapped_targets
    assert "attendance_table" in mapped_targets
    assert "consent_students_table" in mapped_targets


def test_no_fabricated_field_values(pipeline_execution):
    """Test 11: No fabricated or invented field values (signatures, fake attendance)."""
    out_doc = pipeline_execution["out_doc"]
    t1 = out_doc.tables[1]

    # Verify no fake attendance date entries or fake task descriptions
    for r in t1.rows[1:]:
        assert r.cells[0].text.strip() == ""
        assert r.cells[2].text.strip() == ""

    # Verify no fake examiner signatures
    t2 = out_doc.tables[2]
    assert "Name:" in t2.rows[1].cells[0].text
    assert t2.rows[1].cells[1].text.strip() == "" or t2.rows[1].cells[1].text.strip() == "Name:"


def test_unresolved_mappings_diagnostics(pipeline_execution):
    """Test 12: Unresolved mappings produce explicit diagnostics."""
    res = pipeline_execution["result"]
    assert res.structural_report is not None
    assert res.structural_report.is_valid is True
    assert res.structural_report.unresolved_mappings_count > 0

    # Ensure unmapped elements are recorded in manifest
    unmapped = [m for m in res.manifest.mappings if m.status == "unmapped"]
    assert len(unmapped) > 0
    assert any("No template slot" in m.message for m in unmapped)


def test_existing_report_formatting_compatibility():
    """Test 13: Existing report-formatting tests continue to pass (flow generation)."""
    pipeline = DocumentFormattingPipeline(output_dir=tempfile.mkdtemp())
    source = FLOW_SOURCE_PATH
    template = FLOW_TEMPLATE_PATH

    res = pipeline.run(source_path=source, template_path=template, job_id="compat_check")
    assert res.status == JobState.COMPLETED
    assert res.output_docx_path is not None

    doc = docx.Document(res.output_docx_path)
    assert len(doc.paragraphs) > 100
    assert len(doc.sections) >= 1
    # Check that template left margin of 1.25" was adopted
    assert round(doc.sections[0].left_margin.inches, 2) == 1.25
