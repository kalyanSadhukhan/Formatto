import json
import os
import pytest
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.domain.schemas.data_contracts import TemplateConstraint
from app.templates.template_extractor import TemplateExtractor

FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "fixtures", "flow_report")
SAMPLE_PATH = os.path.join(FIXTURES_DIR, "sample.docx")
if not os.path.exists(SAMPLE_PATH):
    SAMPLE_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(__file__)))), "sample.docx")


@pytest.fixture
def complex_template_docx(tmp_path):
    """Create a rich synthetic template DOCX fixture covering Phase 3A & 3B targets."""
    doc_path = tmp_path / "synthetic_template.docx"
    doc = Document()

    # --- Section 0: A4 Portrait, custom margins, title cover page ---
    sec0 = doc.sections[0]
    sec0.page_width = Inches(8.27)  # A4 width (~595.4 pt)
    sec0.page_height = Inches(11.69)  # A4 height (~841.7 pt)
    sec0.orientation = WD_ORIENT.PORTRAIT
    sec0.top_margin = Inches(1.5)
    sec0.bottom_margin = Inches(1.0)
    sec0.left_margin = Inches(1.2)
    sec0.right_margin = Inches(1.2)

    # Cover page title
    p_title = doc.add_paragraph("Synthetic Template Title", style="Title")

    # Add base style: "CompanyPolicy" based on "Normal"
    styles = doc.styles
    custom_style = styles.add_style("CompanyPolicy", WD_STYLE_TYPE.PARAGRAPH)
    custom_style.base_style = styles["Normal"]
    custom_style.font.name = "Georgia"
    custom_style.font.size = Pt(13)
    custom_style.font.bold = True
    custom_style.paragraph_format.space_before = Pt(6)
    custom_style.paragraph_format.space_after = Pt(12)
    custom_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Add child style: "SubPolicy" based on "CompanyPolicy" (inheriting font.name, overriding size and bold)
    sub_style = styles.add_style("SubPolicy", WD_STYLE_TYPE.PARAGRAPH)
    sub_style.base_style = custom_style
    sub_style.font.size = Pt(11)
    sub_style.font.bold = False

    # Add custom heading style: "ExecutiveHeading" with outline level
    exec_heading = styles.add_style("ExecutiveHeading", WD_STYLE_TYPE.PARAGRAPH)
    exec_heading.base_style = styles["Heading 1"] if "Heading 1" in styles else styles["Normal"]
    exec_heading.font.name = "Arial"
    exec_heading.font.size = Pt(16)
    exec_heading.font.bold = True
    # Attach outlineLvl = 1 (Level 2)
    pPr = exec_heading._element.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "1")
    pPr.append(outline)

    # --- Section 1: Landscape section with footer page number ---
    sec1 = doc.add_section(WD_SECTION.NEW_PAGE)
    sec1.page_width = Inches(11.0)
    sec1.page_height = Inches(8.5)
    sec1.orientation = WD_ORIENT.LANDSCAPE
    sec1.top_margin = Inches(0.75)
    sec1.bottom_margin = Inches(0.75)
    sec1.left_margin = Inches(0.75)
    sec1.right_margin = Inches(0.75)

    # Header with text
    hdr = sec1.header
    hdr_p = hdr.paragraphs[0]
    hdr_p.text = "Confidential Report Header"

    # Footer with PAGE number field
    ftr = sec1.footer
    ftr_p = ftr.paragraphs[0]
    ftr_p.text = "Page "
    run = ftr_p.add_run()
    fldSimple = OxmlElement("w:fldSimple")
    fldSimple.set(qn("w:instr"), "PAGE")
    run._r.append(fldSimple)

    # Table with borders
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Header 1"
    table.cell(0, 1).text = "Header 2"
    table.cell(1, 0).text = "Data 1"
    table.cell(1, 1).text = "Data 2"

    # Numbered / bullet list paragraphs
    doc.add_paragraph("Bullet Item", style="List Bullet")
    doc.add_paragraph("Numbered Item", style="List Number")

    # Add paragraph with direct formatting override (manual bold and spacing)
    p_override = doc.add_paragraph("Direct override paragraph", style="Normal")
    p_override.paragraph_format.space_after = Pt(24)
    run_override = p_override.runs[0]
    run_override.font.bold = True
    run_override.font.size = Pt(18)

    # Add an unsupported floating drawing anchor
    drawing_p = doc.add_paragraph()
    r = drawing_p.add_run()
    drawing = OxmlElement("w:drawing")
    anchor = OxmlElement("wp:anchor")
    anchor.set(qn("wp:distT"), "0")
    drawing.append(anchor)
    r._r.append(drawing)

    # Add an unsupported legacy VML shape
    vml_shape = OxmlElement("v:shape")
    vml_shape.set(qn("w:name"), "vmlShape1")
    r._r.append(vml_shape)

    doc.save(str(doc_path))
    return str(doc_path)


def test_synthetic_template_page_setup(complex_template_docx):
    """Test extraction of A4, Landscape, margins, and page sizes."""
    extractor = TemplateExtractor(complex_template_docx)
    profile = extractor.extract_template()

    sections = profile["sections"]
    assert len(sections) == 2

    # Section 0: A4 Portrait
    sec0 = sections[0]
    assert sec0["page_size_name"] == "A4"
    assert sec0["orientation"] == "portrait"
    assert sec0["margins_pt"]["top"] == 108.0  # 1.5 in * 72
    assert sec0["margins_pt"]["bottom"] == 72.0
    assert sec0["margins_pt"]["left"] == 86.4
    assert sec0["margins_pt"]["right"] == 86.4

    # Section 1: Landscape Letter
    sec1 = sections[1]
    assert sec1["orientation"] == "landscape"
    assert sec1["page_size_name"] == "Letter"
    assert sec1["margins_pt"]["top"] == 54.0  # 0.75 in * 72
    assert sec1["has_header"] is True
    assert sec1["header_text"] == "Confidential Report Header"
    assert sec1["has_page_number"] is True


def test_synthetic_template_styles_and_inheritance(complex_template_docx):
    """Test custom style extraction, outline levels, and inheritance resolution."""
    extractor = TemplateExtractor(complex_template_docx)
    profile = extractor.extract_template()
    styles = profile["styles"]

    assert "CompanyPolicy" in styles
    policy = styles["CompanyPolicy"]
    assert policy["is_custom"] is True
    assert policy["based_on"] == "Normal"
    assert policy["raw_declared_properties"]["font_name"] == "Georgia"
    assert policy["raw_declared_properties"]["font_size_pt"] == 13.0
    assert policy["raw_declared_properties"]["bold"] is True

    # SubPolicy inheritance resolution
    assert "SubPolicy" in styles
    sub = styles["SubPolicy"]
    assert sub["based_on"] == "CompanyPolicy"
    assert "font_name" not in sub["raw_declared_properties"]  # not declared in child
    assert sub["inherited_properties"]["font_name"] == "Georgia"  # inherited from parent
    assert sub["property_lineage"]["font_name"] == ["docDefaults", "CompanyPolicy"]
    assert sub["resolved_effective_properties"]["font_name"] == "Georgia"
    assert sub["raw_declared_properties"]["font_size_pt"] == 11.0
    assert sub["resolved_effective_properties"]["font_size_pt"] == 11.0

    # ExecutiveHeading outline level
    assert "ExecutiveHeading" in styles
    exec_h = styles["ExecutiveHeading"]
    assert exec_h["raw_declared_properties"]["outline_level"] == 2


def test_synthetic_template_constraints_provenance(complex_template_docx):
    """Test that generated constraints have valid categories, provenances, and evidence."""
    extractor = TemplateExtractor(complex_template_docx)
    profile = extractor.extract_template()
    constraints = profile["constraints"]

    assert len(constraints) > 0
    valid_categories = {"page", "typography", "paragraph", "structure", "object", "header_footer"}
    valid_provenances = {"EXTRACTED", "CONFIGURED", "INFERRED", "UNSUPPORTED"}

    for c in constraints:
        assert c["category"] in valid_categories
        assert c["provenance"] in valid_provenances
        assert isinstance(c["confidence"], float)
        assert isinstance(c["enforceable"], bool)
        assert "evidence" in c


def test_section_count_and_final_section_properties_sample_docx():
    """
    Verify WordprocessingML section structure on sample.docx:
    - 10 section break paragraphs (w:p/w:pPr/w:sectPr)
    - 1 body section property (w:body/w:sectPr)
    - Exactly 11 total document sections
    - Association of header/footer references with correct sections
    """
    sample_path = SAMPLE_PATH
    if not os.path.exists(sample_path):
        pytest.skip("sample.docx not found in workspace")

    extractor = TemplateExtractor(sample_path)
    profile = extractor.extract_template()
    sections = profile["sections"]

    assert len(sections) == 11

    # Section 0: Title/cover page has no header and no footer
    s0 = sections[0]
    assert s0["section_index"] == 0
    assert s0["break_type"] == "continuous"
    assert s0["has_header"] is False
    assert s0["has_footer"] is False
    assert s0["header_rids"] == []
    assert s0["footer_rids"] == []
    assert s0["has_page_number"] is False

    # Section 1: Explicit footer reference rId8 containing page number
    s1 = sections[1]
    assert s1["section_index"] == 1
    assert s1["break_type"] == "nextPage"
    assert s1["has_footer"] is True
    assert s1["footer_rids"] == ["rId8"]
    assert s1["is_footer_linked"] is False
    assert s1["has_page_number"] is True

    # Section 2: Linked to section 1, inherits footer and page number
    s2 = sections[2]
    assert s2["section_index"] == 2
    assert s2["is_footer_linked"] is True
    assert s2["has_footer"] is True
    assert s2["footer_rids"] == []  # no explicit reference in sectPr
    assert s2["has_page_number"] is True

    # Final Section (10): Captured from w:body/w:sectPr
    s10 = sections[10]
    assert s10["section_index"] == 10
    assert s10["page_size_name"] == "Letter"
    assert s10["orientation"] == "portrait"
    assert s10["margins_pt"]["top"] > 0
    assert s10["margins_pt"]["bottom"] > 0


def test_section_specific_orientation_and_margins(tmp_path):
    """Verify distinct page sizes, orientations, and margins across multiple sections."""
    doc_path = tmp_path / "multi_section.docx"
    doc = Document()

    # Section 0: Portrait A4
    s0 = doc.sections[0]
    s0.page_width = Inches(8.27)
    s0.page_height = Inches(11.69)
    s0.orientation = WD_ORIENT.PORTRAIT
    s0.top_margin = Inches(1.5)
    doc.add_paragraph("Section 0 content")

    # Section 1: Landscape Letter
    s1 = doc.add_section(WD_SECTION.NEW_PAGE)
    s1.page_width = Inches(11.0)
    s1.page_height = Inches(8.5)
    s1.orientation = WD_ORIENT.LANDSCAPE
    s1.top_margin = Inches(0.75)
    doc.add_paragraph("Section 1 content")

    # Section 2: Portrait Legal
    s2 = doc.add_section(WD_SECTION.NEW_PAGE)
    s2.page_width = Inches(8.5)
    s2.page_height = Inches(14.0)
    s2.orientation = WD_ORIENT.PORTRAIT
    s2.top_margin = Inches(1.0)
    doc.add_paragraph("Section 2 content")

    doc.save(str(doc_path))

    extractor = TemplateExtractor(str(doc_path))
    profile = extractor.extract_template()
    secs = profile["sections"]

    assert len(secs) == 3
    assert secs[0]["page_size_name"] == "A4"
    assert secs[0]["orientation"] == "portrait"
    assert secs[0]["margins_pt"]["top"] == 108.0

    assert secs[1]["page_size_name"] == "Letter"
    assert secs[1]["orientation"] == "landscape"
    assert secs[1]["margins_pt"]["top"] == 54.0

    assert secs[2]["page_size_name"] == "Legal"
    assert secs[2]["orientation"] == "portrait"
    assert secs[2]["margins_pt"]["top"] == 72.0


def test_direct_formatting_overrides_detection(complex_template_docx):
    """Verify detection of direct paragraph and run formatting overrides in template body."""
    extractor = TemplateExtractor(complex_template_docx)
    profile = extractor.extract_template()
    overrides = profile["direct_formatting_overrides"]

    assert len(overrides) >= 1
    target = next(o for o in overrides if o["style_id"] == "Normal" and o["text_snippet"] == "Direct override paragraph")
    assert target["paragraph_overrides"]["space_after_pt"] == 24.0
    assert target["run_overrides"]["bold"] is True
    assert target["run_overrides"]["font_size_pt"] == 18.0


def test_unsupported_properties_detection_and_representation(complex_template_docx):
    """Verify detection of floating drawings, VML shapes, and unsupported constraints."""
    extractor = TemplateExtractor(complex_template_docx)
    profile = extractor.extract_template()
    unsupported = profile["unsupported_features"]

    features_detected = {u["feature"] for u in unsupported}
    assert "floating_anchored_drawing" in features_detected
    assert "legacy_vml_shape" in features_detected

    # Verify corresponding UNSUPPORTED constraints exist
    unsupported_constraints = [c for c in profile["constraints"] if c["provenance"] == "UNSUPPORTED"]
    assert len(unsupported_constraints) >= 2
    for uc in unsupported_constraints:
        assert uc["enforceable"] is False
        assert uc["confidence"] == 0.0
        assert uc["expected_value"] == "UNSUPPORTED"


def test_constraint_provenance_and_enforceability_rules(complex_template_docx):
    """
    Verify enforceability contract:
    - EXTRACTED constraints: enforceable = True, confidence = 1.0
    - INFERRED constraints: enforceable = False, confidence = 0.85 (never hard requirements)
    - CONFIGURED constraints: enforceable = False, confidence <= 0.5
    - UNSUPPORTED constraints: enforceable = False, confidence = 0.0
    """
    extractor = TemplateExtractor(complex_template_docx)
    profile = extractor.extract_template()
    constraints = profile["constraints"]

    extracted = [c for c in constraints if c["provenance"] == "EXTRACTED"]
    inferred = [c for c in constraints if c["provenance"] == "INFERRED"]
    unsupported = [c for c in constraints if c["provenance"] == "UNSUPPORTED"]

    assert len(extracted) > 0
    assert len(inferred) > 0
    assert len(unsupported) > 0

    for c in extracted:
        assert c["enforceable"] is True, f"Constraint {c['id']} is EXTRACTED but not enforceable!"
        assert c["confidence"] == 1.0

    for c in inferred:
        assert c["enforceable"] is False, f"Inferred constraint {c['id']} was erroneously marked enforceable!"
        assert c["confidence"] == 0.85

    for c in unsupported:
        assert c["enforceable"] is False
        assert c["confidence"] == 0.0


def test_constraint_schema_pydantic_validation(complex_template_docx):
    """Verify that every constraint parses strictly against the TemplateConstraint Pydantic schema."""
    extractor = TemplateExtractor(complex_template_docx)
    profile = extractor.extract_template()

    for c in profile["constraints"]:
        # Will raise ValidationError if any field violates contract
        validated = TemplateConstraint(**c)
        assert validated.id.startswith("tc_")
        assert validated.category in {"page", "typography", "paragraph", "structure", "object", "header_footer"}
        assert validated.provenance in {"EXTRACTED", "CONFIGURED", "INFERRED", "UNSUPPORTED"}
        assert validated.evidence is not None


def test_real_template_integration_sample_docx():
    """Integration test against actual academic report sample.docx."""
    sample_path = SAMPLE_PATH
    if not os.path.exists(sample_path):
        pytest.skip("sample.docx not found in workspace")

    extractor = TemplateExtractor(sample_path)
    profile = extractor.extract_template()

    assert profile["template_id"] == "bd973da4"
    assert len(profile["sections"]) == 11
    assert len(profile["constraints"]) >= 80

    # Section 0 has no page number; later sections do
    assert profile["sections"][0]["has_page_number"] is False
    assert any(s["has_page_number"] for s in profile["sections"][1:])

    # Validate that all constraints satisfy schema
    for c in profile["constraints"]:
        TemplateConstraint(**c)


def test_template_idempotency_repeatability(complex_template_docx):
    """Test that running template extraction multiple times produces identical output."""
    runs = []
    for _ in range(3):
        ext = TemplateExtractor(complex_template_docx)
        runs.append(ext.extract_template())

    assert runs[0] == runs[1]
    assert runs[1] == runs[2]


def test_deterministic_serialization_byte_identical():
    """Verify that repeat extraction produces byte-for-byte identical JSON serialization."""
    sample_path = SAMPLE_PATH
    if not os.path.exists(sample_path):
        pytest.skip("sample.docx not found in workspace")

    ext1 = TemplateExtractor(sample_path)
    prof1 = ext1.extract_template()
    json1 = json.dumps(prof1, sort_keys=True, indent=2)

    ext2 = TemplateExtractor(sample_path)
    prof2 = ext2.extract_template()
    json2 = json.dumps(prof2, sort_keys=True, indent=2)

    assert json1 == json2


def test_corrupt_template_failure(tmp_path):
    """Test that corrupt or missing template raises clean ValueError."""
    corrupt_file = tmp_path / "bad_template.docx"
    with open(corrupt_file, "wb") as f:
        f.write(b"CORRUPT_BYTES")

    with pytest.raises(ValueError):
        TemplateExtractor(str(corrupt_file))
