import io
import os
import tempfile
import zipfile
import pytest
import docx
from fastapi.testclient import TestClient

from app.domain.schemas.data_contracts import JobState
from app.extraction.docx_extractor import DocxExtractor
from app.main import app
from app.pipeline.orchestrator import DocumentFormattingPipeline

FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "fixtures")
PROJECT_DIARY_DIR = os.path.join(FIXTURES_DIR, "project_diary")
FLOW_REPORT_DIR = os.path.join(FIXTURES_DIR, "flow_report")

FAILING_SOURCE = os.path.join(PROJECT_DIARY_DIR, "test2_raw_source.docx")
FAILING_TEMPLATE = os.path.join(PROJECT_DIARY_DIR, "cis_diary_26_odd_template.docx")
SAMPLE_SOURCE = os.path.join(FLOW_REPORT_DIR, "sample.docx")
REFERENCE_TEMPLATE = os.path.join(FLOW_REPORT_DIR, "reference_template.docx")


def test_reproduction_and_fix_of_table_merge_index_error():
    """
    Task 1 & 2: Verify that extracting the exact failing source document
    (test2 raw.docx with merged cells and asymmetric row cell counts)
    does not raise IndexError: list index out of range and extracts all tables.
    """
    assert os.path.exists(FAILING_SOURCE), f"Fixture not found: {FAILING_SOURCE}"
    extractor = DocxExtractor(FAILING_SOURCE)
    elements = extractor.extract()

    assert len(elements) > 0
    tables = [e for e in elements if e.type == "table"]
    # test2_raw_source.docx contains 13 tables
    assert len(tables) == 13

    # Table at index 7 has asymmetric rows (row 0 has 4 cells with gridSpan=3, row 1 has 6 cells with vMerge)
    t7 = tables[7]
    assert t7.metadata["has_merged_cells"] is True
    assert len(t7.metadata["merged_cells_info"]) == 8


def test_pipeline_failing_combination_end_to_end():
    """
    Verify complete pipeline execution with the user's failing document pair
    (test2_raw_source.docx + cis_diary_26_odd_template.docx).
    """
    out_dir = tempfile.mkdtemp(prefix="test_regress_exec_")
    pipeline = DocumentFormattingPipeline(output_dir=out_dir)
    res = pipeline.run(
        source_path=FAILING_SOURCE,
        template_path=FAILING_TEMPLATE,
        job_id="test_regress_comb",
    )

    assert res.status == JobState.COMPLETED
    assert res.output_docx_path is not None
    assert os.path.exists(res.output_docx_path)
    assert res.structural_report is not None
    assert res.structural_report.is_valid is True

    # Validate generated docx opens and has 13 tables matching the template
    doc = docx.Document(res.output_docx_path)
    assert len(doc.tables) == 13


def test_synthetic_merged_cells_with_asymmetric_counts_and_empty_paragraphs():
    """
    Test extraction of synthetic table with empty paragraphs and asymmetric row cell counts.
    Row 0: 2 physical cells, cell 1 has gridSpan=2 (covers grid cols 1 and 2)
    Row 1: 3 physical cells, cell 2 has vMerge="restart"
    Row 2: 3 physical cells, cell 2 has vMerge="continue"
    """
    doc = docx.Document()
    # Add empty paragraph
    doc.add_paragraph("")
    doc.add_paragraph("Synthetic Table Heading")

    tbl = doc.add_table(rows=3, cols=3)
    # Set text on row 0
    tbl.cell(0, 0).text = "Header 0"
    tbl.cell(0, 1).text = "Merged Header 1-2"
    # Merge row 0 cells 1 and 2 horizontally
    tbl.cell(0, 1).merge(tbl.cell(0, 2))

    # Set vMerge on column 2 of row 1 and 2
    tbl.cell(1, 0).text = "R1C0"
    tbl.cell(1, 1).text = "R1C1"
    tbl.cell(1, 2).text = "VMerge Origin"
    tbl.cell(2, 0).text = "R2C0"
    tbl.cell(2, 1).text = "R2C1"
    tbl.cell(2, 2).text = ""

    # Merge cell (1, 2) and (2, 2) vertically
    tbl.cell(1, 2).merge(tbl.cell(2, 2))

    # Add empty paragraph at end
    doc.add_paragraph("")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name
    doc.save(tmp_path)

    try:
        extractor = DocxExtractor(tmp_path)
        elements = extractor.extract()
        table_elements = [e for e in elements if e.type == "table"]
        assert len(table_elements) == 1
        t_meta = table_elements[0].metadata
        assert t_meta["has_merged_cells"] is True
        assert t_meta["row_count"] == 3
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)


def test_flow_report_workflow_end_to_end():
    """
    Verify existing sample.docx + reference_template.docx flow report workflow.
    """
    assert os.path.exists(SAMPLE_SOURCE), f"Fixture not found: {SAMPLE_SOURCE}"
    assert os.path.exists(REFERENCE_TEMPLATE), f"Fixture not found: {REFERENCE_TEMPLATE}"

    out_dir = tempfile.mkdtemp(prefix="test_flow_exec_")
    pipeline = DocumentFormattingPipeline(output_dir=out_dir)
    res = pipeline.run(
        source_path=SAMPLE_SOURCE,
        template_path=REFERENCE_TEMPLATE,
        job_id="test_flow_regress",
    )

    assert res.status == JobState.COMPLETED
    assert res.output_docx_path is not None
    assert os.path.exists(res.output_docx_path)
    doc = docx.Document(res.output_docx_path)
    assert len(doc.paragraphs) > 0


def test_api_invalid_file_extension():
    """API should return HTTP 400 for non-.docx files."""
    client = TestClient(app)
    response = client.post(
        "/api/jobs",
        files={
            "source_file": ("test.txt", io.BytesIO(b"not docx content"), "text/plain"),
            "template_file": ("template.docx", io.BytesIO(b"fake docx"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        },
    )
    assert response.status_code == 400
    data = response.json()
    assert data["detail"]["error_code"] == "INVALID_UPLOAD"
    assert "Only .docx documents are supported" in data["detail"]["error"]


def test_api_empty_file_upload():
    """API should return HTTP 400 for empty (0 bytes) files."""
    client = TestClient(app)
    response = client.post(
        "/api/jobs",
        files={
            "source_file": ("empty.docx", io.BytesIO(b""), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "template_file": ("template.docx", io.BytesIO(b"fake docx"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        },
    )
    assert response.status_code == 400
    data = response.json()
    assert data["detail"]["error_code"] == "INVALID_UPLOAD"
    assert "empty (0 bytes)" in data["detail"]["error"]


def test_api_corrupted_file_upload():
    """API should return HTTP 400 for corrupted non-zip files named .docx."""
    client = TestClient(app)
    response = client.post(
        "/api/jobs",
        files={
            "source_file": ("corrupted.docx", io.BytesIO(b"this is definitely not a zip file"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "template_file": ("template.docx", io.BytesIO(b"neither is this"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        },
    )
    assert response.status_code == 400
    data = response.json()
    assert data["detail"]["error_code"] == "UNSUPPORTED_FILE"
    assert "corrupted or not a valid DOCX package" in data["detail"]["error"]


def test_api_valid_job_creation_and_download():
    """API should accept valid documents, generate output, and allow download."""
    client = TestClient(app)
    with open(FAILING_SOURCE, "rb") as sf, open(FAILING_TEMPLATE, "rb") as tf:
        response = client.post(
            "/api/jobs",
            files={
                "source_file": ("test2_raw_source.docx", sf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
                "template_file": ("cis_diary_26_odd_template.docx", tf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            },
        )
    assert response.status_code == 200
    data = response.json()
    job_id = data["job_id"]
    assert data["status"] == "CREATED"

    # Status endpoint
    status_resp = client.get(f"/api/jobs/{job_id}/status")
    assert status_resp.status_code == 200
    assert status_resp.json()["status"] == "COMPLETED"

    # Result endpoint
    res_resp = client.get(f"/api/jobs/{job_id}/result")
    assert res_resp.status_code == 200
    res_data = res_resp.json()
    assert res_data["status"] == "COMPLETED"
    assert res_data["metadata"]["structural_validation"]["is_valid"] is True

    # Download endpoint
    dl_resp = client.get(f"/api/jobs/{job_id}/download")
    assert dl_resp.status_code == 200
    assert len(dl_resp.content) > 10000
    assert zipfile.is_zipfile(io.BytesIO(dl_resp.content))
